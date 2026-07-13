import re
import os
import pypdf
import docx
import requests
import json

# Master list of commonly tracked skills across domains
COMMON_SKILLS = [
    # Programming Languages
    "python", "javascript", "typescript", "java", "c++", "c#", "php", "ruby", "go", "golang",
    "rust", "swift", "kotlin", "sql", "html", "css", "html5", "css3", "r", "bash", "shell",
    
    # Frameworks & Libraries
    "react", "react.js", "reactjs", "vue", "vue.js", "angular", "node.js", "nodejs", "express",
    "flask", "django", "fastapi", "spring", "spring boot", "bootstrap", "tailwind", "jquery",
    "next.js", "nuxt", "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "keras",
    
    # Databases & Cloud
    "mysql", "postgresql", "sqlite", "mongodb", "redis", "oracle", "sql server", "dynamodb",
    "aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "git", "github", "gitlab",
    "ci/cd", "jenkins", "terraform", "linux", "nginx", "apache",
    
    # Developer Concepts & Methodologies
    "rest api", "graphql", "microservices", "oop", "object-oriented programming", "data structures",
    "algorithms", "agile", "scrum", "kanban", "unit testing", "system design", "devops",
    
    # Soft & Professional Skills
    "problem solving", "communication", "leadership", "teamwork", "time management",
    "project management", "critical thinking", "analytical skills", "collaboration"
]

def extract_text_from_pdf(filepath):
    text = ""
    try:
        reader = pypdf.PdfReader(filepath)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        print(f"Error reading PDF {filepath}: {e}")
    return text

def extract_text_from_docx(filepath):
    text = ""
    try:
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
    except Exception as e:
        print(f"Error reading DOCX {filepath}: {e}")
    return text

def extract_text_from_txt(filepath):
    text = ""
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    except Exception as e:
        print(f"Error reading TXT {filepath}: {e}")
    return text

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(filepath)
    elif ext == '.docx':
        return extract_text_from_docx(filepath)
    elif ext == '.txt':
        return extract_text_from_txt(filepath)
    return ""

def extract_skills_from_text(text):
    text_clean = text.lower()
    found_skills = set()
    
    for skill in COMMON_SKILLS:
        # Match as whole word/phrase
        pattern = r'\b' + re.escape(skill) + r'\b'
        if re.search(pattern, text_clean):
            found_skills.add(skill.capitalize() if len(skill) > 3 else skill.upper())
            
    return sorted(list(found_skills))

def analyze_resume_ats(resume_text, job_description, job_title="Job Candidate"):
    res_text_clean = resume_text.lower()
    jd_text_clean = job_description.lower()
    
    # 1. Skill Matching
    jd_skills = extract_skills_from_text(job_description)
    resume_skills = extract_skills_from_text(resume_text)
    
    # If no specific predefined skills found in JD, extract custom n-grams/words from JD
    if not jd_skills:
        jd_words = set(re.findall(r'\b[a-zA-Z]{3,}\b', jd_text_clean))
        stop_words = {'and', 'the', 'for', 'with', 'you', 'will', 'are', 'this', 'that', 'have', 'from', 'your', 'work'}
        jd_skills = sorted(list(jd_words - stop_words))[:10]
        
    matched_skills = [s for s in jd_skills if any(re.search(r'\b' + re.escape(s.lower()) + r'\b', res_text_clean) for _ in [1])]
    missing_skills = [s for s in jd_skills if s not in matched_skills]
    
    if jd_skills:
        skills_score = int((len(matched_skills) / len(jd_skills)) * 100)
    else:
        skills_score = 70  # Default fallback
        
    skills_score = min(100, max(0, skills_score))

    # 2. ATS Formatting & Structure Checks
    formatting_feedback = []
    ats_formatting_score = 100
    
    word_count = len(re.findall(r'\w+', resume_text))
    if word_count < 150:
        formatting_feedback.append({
            "status": "warning",
            "title": "Low Word Count",
            "message": f"Your resume has only ~{word_count} words. Aim for 300 to 800 words for adequate detail."
        })
        ats_formatting_score -= 15
    elif word_count > 1200:
        formatting_feedback.append({
            "status": "warning",
            "title": "High Word Count",
            "message": f"Your resume is quite long (~{word_count} words). Concise resumes perform better in ATS systems."
        })
        ats_formatting_score -= 10
    else:
        formatting_feedback.append({
            "status": "success",
            "title": "Optimal Word Count",
            "message": f"Word count ({word_count} words) is well-suited for ATS scanners."
        })

    # Contact Info Check
    email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', resume_text)
    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', resume_text)
    
    if email_match:
        formatting_feedback.append({
            "status": "success",
            "title": "Contact Information",
            "message": f"Email address found ({email_match.group(0)})."
        })
    else:
        formatting_feedback.append({
            "status": "danger",
            "title": "Missing Email",
            "message": "No valid email address detected. Make sure your email is explicitly listed in plain text."
        })
        ats_formatting_score -= 20

    if phone_match:
        formatting_feedback.append({
            "status": "success",
            "title": "Phone Number",
            "message": "Phone number detected."
        })
    else:
        formatting_feedback.append({
            "status": "warning",
            "title": "Missing Phone Number",
            "message": "Could not verify phone number format."
        })
        ats_formatting_score -= 10

    # Essential Sections Check
    sections = {
        "Experience": r'\b(experience|work history|employment|history|professional experience)\b',
        "Education": r'\b(education|academic|qualification|degree)\b',
        "Skills": r'\b(skills|technical skills|technologies|proficiencies)\b',
        "Projects": r'\b(projects|key projects|portfolio)\b'
    }
    
    found_sections = []
    missing_sections = []
    
    for sec_name, sec_pattern in sections.items():
        if re.search(sec_pattern, res_text_clean):
            found_sections.append(sec_name)
        else:
            missing_sections.append(sec_name)
            
    if missing_sections:
        formatting_feedback.append({
            "status": "warning",
            "title": "Section Headings",
            "message": f"Consider adding standard headings for: {', '.join(missing_sections)}."
        })
        ats_formatting_score -= (len(missing_sections) * 5)
    else:
        formatting_feedback.append({
            "status": "success",
            "title": "Section Headings",
            "message": "All essential resume section headers detected (Experience, Education, Skills, Projects)."
        })

    ats_formatting_score = min(100, max(20, ats_formatting_score))

    # Overall Weighted Score: 60% skills match + 40% formatting & readability
    overall_score = int((skills_score * 0.6) + (ats_formatting_score * 0.4))

    return {
        "overall_score": overall_score,
        "skills_score": skills_score,
        "ats_score": ats_formatting_score,
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "formatting_feedback": formatting_feedback
    }

def analyze_with_gemini(api_key, resume_text, job_description, job_title="Candidate"):
    if not api_key:
        # Graceful fallback when API key is missing
        return {
            "cover_letter": f"Dear Hiring Manager,\n\nI am writing to express my strong interest in the {job_title} position. With my background in technology and professional skills, I am confident in my ability to contribute to your team.\n\nThank you for your time and consideration.\n\nSincerely,\n{job_title} Applicant",
            "job_fit_analysis": "Basic analysis run. Configure a valid GEMINI_API_KEY to receive custom AI-driven strengths and weaknesses assessments.",
            "ai_suggestions": [
                {
                    "section_name": "Skills Section",
                    "current_text": "Review your list of professional technical skills.",
                    "suggested_text": "Tailor your skills section to highlight keywords matching the target job description directly.",
                    "reason": "ATS scanners prioritize candidates whose skill headings and lists match the spelling in the job posting exactly."
                }
            ]
        }
    
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={api_key}"
    headers = {
        "Content-Type": "application/json"
    }
    
    prompt = f"""
    Analyze the following resume against the job description for the target job title.
    Target Job Title: {job_title}
    
    Resume Text:
    ---
    {resume_text}
    ---
    
    Job Description:
    ---
    {job_description}
    ---
    
    You must return a JSON object with the following fields:
    1. "cover_letter": A professional, tailored cover letter (about 250-300 words) from the perspective of the candidate for the target job role. Use standard placeholders for personal details if not found.
    2. "job_fit_analysis": A brief paragraph summarizing how well the candidate's resume fits the target job description, noting key strengths.
    3. "ai_suggestions": An array of objects. Each object should represent a specific recommendation to rewrite or optimize a section of the resume. Each object must have these exact keys:
       - "section_name": The name of the section (e.g., Professional Summary, Work Experience, Projects).
       - "current_text": A snippet of text from the user's resume that needs improvement.
       - "suggested_text": The rewritten/improved version of that text, optimized with relevant keywords from the job description.
       - "reason": A short explanation of why this rewrite helps (e.g. highlights leadership, adds SQL keyword).
    
    Ensure the JSON matches the schema exactly and contains valid JSON content.
    """
    
    payload = {
        "contents": [{
            "parts": [{
                "text": prompt
            }]
        }],
        "generationConfig": {
            "responseMimeType": "application/json"
        }
    }
    
    try:
        response = requests.post(url, headers=headers, json=payload, timeout=30)
        if response.status_code == 200:
            res_data = response.json()
            text_response = res_data['candidates'][0]['content']['parts'][0]['text']
            return json.loads(text_response)
        else:
            print(f"Gemini API returned error code {response.status_code}: {response.text}")
    except Exception as e:
        print(f"Error calling Gemini API: {e}")
        
    # Default fallback on failure
    return {
        "cover_letter": f"Dear Hiring Manager,\n\nI am writing to express my strong interest in the {job_title} position. With my background in technology and professional skills, I am confident in my ability to contribute to your team.\n\nThank you for your time and consideration.\n\nSincerely,\n{job_title} Applicant",
        "job_fit_analysis": "The Gemini AI analysis was unable to complete successfully. Please verify your internet connection or API Key credentials.",
        "ai_suggestions": [
            {
                "section_name": "Skills Section",
                "current_text": "Review your list of professional technical skills.",
                "suggested_text": "Tailor your skills section to highlight keywords matching the target job description directly.",
                "reason": "ATS scanners prioritize candidates whose skill headings and lists match the spelling in the job posting exactly."
            }
        ]
    }
